# SPDX-License-Identifier: Apache-2.0 OR MIT
"""Export a `.mdtbl` to `.xlsx` — deliberately OUTSIDE `reference/`.

`reference/` is standard-library-only because a dependency there is a
dependency every independent implementation inherits, and an independent
implementation of a table format must not be required to read xlsx. So the
exporter lives here, ships as the `rowspec[xlsx]` extra, and the boundary is
held by a test that fails if `reference/` ever imports it — `tests/test_boundary.py`
— rather than by anyone remembering this paragraph.

The direction is one-way on purpose: this package imports `rowspec`, and
nothing in `rowspec` may import this package.

What it writes (kindspec/rowspec#35): one sheet holding one spreadsheet TABLE
named `rowspec`, whose computed columns are live formulas over structured
references (`rowspec[[#This Row],[qty]]`), and the declared aggregates below
the grid as `SUBTOTAL` forms over the table's columns. A `.mdtbl` addresses by
name and so does a structured reference, which is why the export is worth
having, and why nothing here ever falls back to `A1`: there is no coordinate
in any formula this module writes, and `tests/test_xlsx_export.py` greps for
one.

**The formulas carry this format's rules, not the spreadsheet's.** A
spreadsheet reads a blank operand as `0`, skips text inside `SUM`, and answers
`0` for the `MIN` of nothing — three plausible numbers where SPEC.md §7 and §8
require an error. So every operand is guarded (`IF(ISNUMBER(x),x,NA())`), and
every aggregate refuses a column holding text or an error, so a `#REF!` in the
source arrives as a spreadsheet error value — loud, and propagating through
everything downstream of it — never as a blank, a zero or a cached number. The
error's NAME does not survive: a spreadsheet error cannot carry `(qty)`.

**What a spreadsheet cannot say, the export refuses, rather than
approximating** (`NotExpressible`, and nothing is written):

- `cumulative`, `prior` and `delta` need the preceding row in a declared order,
  and a spreadsheet can only say "the row above" with a position;
- a `where` predicate, on a declaration or in a per-row group aggregate: the
  spreadsheet functions that take a criterion match case-insensitively and read
  `*`, `?` and a leading `<` or `=` as syntax, where §4.2 rule 6 compares exact
  text, and `SUBTOTAL` takes no criterion at all;
- an `if` comparing a column against a non-empty string when that column holds
  a number: the spreadsheet stores the number and forgets the spelling the
  comparison reads (`3.0` against `"3"`);
- two column names that differ only by case, which a spreadsheet table header
  cannot hold;
- a numeric literal with more significant digits than a spreadsheet formula
  keeps, and a table with no rows, which a spreadsheet table cannot be;
- a character XML 1.0 has no place for, such as a form feed in a stored cell
  or a string literal: every C0 control but tab, LF and CR, and U+FFFE and
  U+FFFF. §4.1 allows them, and the XML an `.xlsx` is made of cannot hold
  them.

**The workbook recalculates under spreadsheet arithmetic, not this format's,
and it can disagree with `rowspec eval` by more than the last digit.** A
spreadsheet compares numbers with a tolerance and snaps a near-equal
subtraction to `0`, so `if(c = 0.3, ...)` over `0.1 + 0.2` takes the other
branch (#88). It also adds left to right where §7 takes the exact sum, so
`1e16 + 1 - 1e16` totals `0` where `rowspec eval` says `1` (#83). The formulas
cannot switch either behaviour off; they belong to the engine.

Two further differences are deliberate and stay:

- the KEY column is written as text, because §9.16 makes a key an identifier,
  so a formula that uses the key as a number (`x = id * 2`) is `NA()` in the
  workbook where `rowspec eval` computes it;
- `sum`, `avg`, `min` and `max` are `SUBTOTAL` forms, which leave out rows a
  spreadsheet user hides or filters, while `count` is `ROWS` and counts them.
  With nothing hidden the two agree.

A stored cell is still never a formula: `_append` forces every string cell to
text, because a stored cell is text under §5 and a spreadsheet application
would otherwise read a leading `=` as a live formula.

Every number cell, stored or computed, is saved with the shortest spelling
that reads back as the same binary64 value (`_Workbook.save`), where openpyxl
alone keeps 16 significant digits (#269).

**`values=True` writes `rowspec eval`'s values instead of formulas**
(kindspec/rowspec#89). Every computed cell and aggregate is the binary64 value
`evaluate` computed, so the numbers are exactly the format's, and none of the
spreadsheet arithmetic above applies. A `#REF!` is written as the spreadsheet
error `#REF!`, an error cell and not text, so a spreadsheet `SUM` over it
answers an error instead of skipping it; its name is lost, as in a formula
export. The workbook holds no formula, so nothing in it recalculates: it is
the table as computed when it was exported.

What the formula export refuses only because no formula says it is exported
in this mode, since `evaluate` has its value: `cumulative`, `prior`, `delta`,
`where` on a column or a declaration, the `if` that compares a number column
with text, the literal of more than 15 significant digits, and a string
literal holding a character XML cannot hold. Three refusals stay, because
they are about the sheet and not about a formula: column names that differ
only by case and a table with no rows, which a spreadsheet table cannot be,
and a character XML cannot hold in a header, a stored cell or the sheet
title.
"""

try:
    from openpyxl import Workbook
    from openpyxl.cell import _writer as _cell_writer
    from openpyxl.workbook.properties import CalcProperties
    from openpyxl.worksheet.table import Table, TableColumn, TableFormula
except ModuleNotFoundError as exc:  # the extra is absent, which is the default
    raise ModuleNotFoundError(
        "xlsx export is an optional extra of rowspec, and it is not installed: "
        "pip install 'rowspec[xlsx]'"
    ) from exc

import math
import re
import threading

from rowspec.table import (
    _GROUP,
    ROWREL,
    Bin,
    If,
    Name,
    Neg,
    Num,
    _ast,
    evaluate,
    names_of,
    num,
    parse,
)

__all__ = ["NotExpressible", "export_file", "to_workbook"]

TABLE = "rowspec"
# A spreadsheet keeps 15 significant decimal digits of a number typed into a
# formula, so a longer literal would be rounded on the way in.
_MAX_LITERAL_DIGITS = 15


class NotExpressible(Exception):
    """The table is valid, and a spreadsheet cannot say what it says."""


# What XML 1.0's `Char` leaves out: every C0 control but tab, LF and CR, the
# surrogates, and U+FFFE and U+FFFF. openpyxl refuses only the C0 controls and
# writes the others into a file no reader can parse.
_NOT_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _expressible(text):
    if _NOT_XML.search(text):
        raise NotExpressible(f"{text!r}: holds a character no spreadsheet file can hold")


def _value(v):
    """Numbers arrive as numbers, everything else as text, blanks as blank.

    A stored cell matching §4.1's number grammar already IS a number to §8's
    evaluator, so writing it into a spreadsheet as text would contradict the
    format's own reading of the same cell — and text that looks like a number
    is the single most annoying thing to be handed in a workbook.

    `num` rather than `float` on purpose: it is the format's grammar, and it is
    deliberately narrower. `1_000`, `1e3` and Arabic-Indic digits are text
    here, exactly as they are under §8.

    A cell whose spelling is a number but whose binary64 value is infinite
    stays text: §8 makes it `#REF!(overflow)` wherever it is used as an
    operand, and a spreadsheet cannot store an infinity.

    This is a VALUE policy and it must not be applied to the key column: §9.16
    makes a key an identifier, not a number. See `to_workbook`.
    """
    try:
        got = num(str(v))
    except ValueError:
        return str(v) or None
    return got if math.isfinite(got) else str(v)


def _append(ws, values):
    """Append one row, then force every string cell to TEXT.

    openpyxl infers a cell's type from its value, so a stored cell of `=1+1`
    lands as a LIVE FORMULA (`data_type == "f"`) and survives save and reload
    as one. Three things are wrong with that. §5 makes a stored cell text, so
    promoting it to a computed cell is not the file the format describes. A
    `=A1` would put a COORDINATE in the output, which is the one thing this
    format does not have. And a `.mdtbl` is a file anyone can open a pull
    request against, so it is formula injection with a review step in front of
    it.

    The formulas this module writes go through `_formula`, never through here.
    """
    for v in values:
        if isinstance(v, str):
            _expressible(v)
    ws.append(values)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _formula(cell, text):
    _expressible(text)
    cell.value = "=" + text
    cell.data_type = "f"


def _static(cell, v):
    """One of `evaluate`'s values, as it is. §4.2 rule 10 gives a computed
    value two kinds, a number and an error, and §7 lets a blank through.
    `#REF!(qty)` becomes the error `#REF!`, never the text, because a
    spreadsheet `SUM` skips text and a `#REF!` must not be skipped."""
    if isinstance(v, str) and v.startswith("#REF!"):
        cell.value = "#REF!"
        cell.data_type = "e"
    elif v == "":
        cell.value = None
    elif isinstance(v, int | float) and not isinstance(v, bool):
        cell.value = v
    else:
        raise TypeError(f"{v!r}: a computed value is a number, a blank or a #REF!")


# openpyxl writes a number cell's `<v>` with `safe_string`, which spells a
# float as `%.16g`: a digit short of binary64, so 0.30000000000000004 is saved
# as 0.3, and -0.0 as `-0`, which reads back as the integer 0 (#269). Bound at
# import, so an openpyxl without the name fails here, not silently at save.
# pyproject pins openpyxl below 3.2 because this is a private module, and
# tests/test_xlsx_export.py fails if a save stops going through it.
_SAFE_STRING = _cell_writer.safe_string


def _exact(value):
    """`float.__repr__` for a finite float: the shortest spelling that reads
    back as the same double. Everything else is openpyxl's own spelling.

    `float.__repr__`, not `repr`: a float subclass may spell itself otherwise,
    and `numpy.float64` is `np.float64(0.1)`, which is not a number in XML."""
    if isinstance(value, float) and math.isfinite(value):
        return float.__repr__(value)
    return _SAFE_STRING(value)


# Held for the whole install, save and restore, so `_Workbook` saves are
# serialized: unguarded, one thread's restore lands in the middle of another's
# save, and the rest of that file is written with 16 digits.
_SAVE_LOCK = threading.Lock()


class _Workbook(Workbook):
    """A workbook whose `save` writes every number cell exactly.

    The override of openpyxl's number-to-string step lasts only for the save,
    so the cells in memory keep their real values. `_Workbook` saves are
    serialized by `_SAVE_LOCK`. A plain openpyxl save running at the same
    moment in another thread also sees the override, so its finite floats get
    the same exact spelling instead of openpyxl's 16 digits.

    Only this `save` installs the override. Calling openpyxl's
    `save_workbook(wb, ...)` directly bypasses it and writes 16 digits.
    """

    def save(self, filename):
        with _SAVE_LOCK:
            _cell_writer.safe_string = _exact
            try:
                super().save(filename)
            finally:
                _cell_writer.safe_string = _SAFE_STRING


def _ref(col):
    return f"{TABLE}[[#This Row],[{col}]]"


def _operand(col):
    # §4.2: a name in an expression asserts the cell holds a number. A blank is
    # `0` to a spreadsheet and `#REF!` here, and text is coerced by some
    # spreadsheets and refused by §8.
    r = _ref(col)
    return f"IF(ISNUMBER({r}),{r},NA())"


def _literal(v, where):
    text = repr(float(v)).upper()
    digits = re.sub(r"E.*", "", text).replace("-", "").replace(".", "").strip("0")
    if len(digits) > _MAX_LITERAL_DIGITS:
        raise NotExpressible(
            f"column {where!r}: the literal {v!r} has more than "
            f"{_MAX_LITERAL_DIGITS} significant digits, and a spreadsheet would "
            f"round it on the way in"
        )
    return text[:-2] if text.endswith(".0") else text


def _translate(node, where, text_cmp):
    """One `expr` tree (§4.2) as a spreadsheet formula, with this format's rules."""
    if isinstance(node, Num):
        return _literal(node.v, where)
    if isinstance(node, Name):
        return _operand(node.id)
    if isinstance(node, Neg):
        return f"-({_translate(node.x, where, text_cmp)})"
    if isinstance(node, Bin):
        left = _translate(node.l, where, text_cmp)
        right = _translate(node.r, where, text_cmp)
        return f"({left}{node.op}{right})"
    if isinstance(node, If):
        c = node.c
        if c.kind == "str":
            text_cmp.append((c.lhs, c.rhs))
            # EXACT, not `=`: a spreadsheet's `=` on text ignores case, and
            # §4.2 rule 6 compares the characters. A blank is the empty string
            # to both, which is the blank test of rule 10.
            test = f'EXACT({_ref(c.lhs)},"{c.rhs}")'
            if c.op == "<>":
                test = f"NOT({test})"
        else:
            rhs = _literal(c.rhs, where) if c.kind == "num" else _operand(c.rhs)
            test = f"{_operand(c.lhs)}{c.op}{rhs}"
        a = _translate(node.a, where, text_cmp)
        b = _translate(node.b, where, text_cmp)
        return f"IF({test},{a},{b})"
    raise NotExpressible(f"column {where!r}: no spreadsheet form for {node!r}")


def _on_cycle(formulas, deps):
    """Computed columns that reach themselves. A spreadsheet answers a circular
    reference with a warning and a stale value; §4.2 rule 9 answers with an
    error, so those columns are written as one."""
    out = set()
    for start in formulas:
        stack, seen = list(deps[start]), set()
        while stack:
            n = stack.pop()
            if n == start:
                out.add(start)
                break
            if n in seen or n not in formulas:
                continue
            seen.add(n)
            stack.extend(deps[n])
    return out


def _column_formulas(cols, formulas, rows, key):
    """Each computed column's spreadsheet formula, or refuse the export."""
    for nm, expr in formulas.items():
        e = expr.strip()
        if any(re.fullmatch(rf"{f}\s*\(\s*\w+\s*\)", e) for f in ROWREL):
            raise NotExpressible(
                f"column {nm!r}: {e!r} needs the preceding row in the declared "
                f"order, and a spreadsheet can only name that row by its position"
            )
        if _GROUP.fullmatch(e):
            raise NotExpressible(
                f"column {nm!r}: {e!r} filters with `where`, and no spreadsheet "
                f"criterion compares exact text the way §4.2 rule 6 does"
            )
    trees = {nm: _ast(expr, nm) for nm, expr in formulas.items()}
    deps = {nm: names_of(t) for nm, t in trees.items()}
    cyclic = _on_cycle(formulas, deps)
    out, text_cmp = {}, []
    for nm, tree in trees.items():
        if nm in cyclic or any(n not in cols for n in deps[nm]):
            # A name the header does not have, or a cycle: §8 makes the whole
            # column an error, decided by the header alone. A structured
            # reference to an absent column is not a formula a spreadsheet
            # will open, so the error is written directly.
            out[nm] = "NA()"
        else:
            out[nm] = _translate(tree, nm, text_cmp)
    for col, lit in text_cmp:
        if lit == "" or col == key:
            continue  # the blank test reads no spelling, and a key is text
        for r in rows:
            v = r.get(col)
            if isinstance(_value(v), float):
                raise NotExpressible(
                    f"column {col!r} is compared with the text {lit!r}, and holds "
                    f"the number {v!r}: a spreadsheet stores the number and not the "
                    f"spelling the comparison reads"
                )
    return out


_SUBTOTAL = {"sum": 109, "avg": 101, "min": 105, "max": 104}


def _aggregate(fn, col, pred, cols, name):
    """A declaration (§7) as a formula over the table's column."""
    if pred:
        raise NotExpressible(
            f"aggregate {name!r}: `where {pred}` has no spreadsheet form; SUBTOTAL "
            f"takes no criterion, and the functions that do match text "
            f"case-insensitively and read `*`, `?`, `<` and `=` as syntax"
        )
    if col not in cols:
        return "NA()"  # §8: #REF!(col)
    rng = f"{TABLE}[{col}]"
    if fn == "count":
        # §7: `count` counts rows and never coerces, and a #REF! actually
        # present in the column still poisons it.
        return f"IF(SUMPRODUCT(--ISERROR({rng}))>0,NA(),ROWS({rng}))"
    # §7: a blank is skipped, and text or an error poisons. SUBTOTAL skips text
    # silently, so a column whose non-empty cells (103) are not all numbers
    # (102) is an error before SUBTOTAL ever sees it.
    form = f"IF(SUBTOTAL(103,{rng})<>SUBTOTAL(102,{rng}),NA(),SUBTOTAL({_SUBTOTAL[fn]},{rng}))"
    if fn in ("min", "max"):
        # A spreadsheet's MIN of nothing is 0; §7's is #REF!(empty).
        form = f"IF(SUBTOTAL(102,{rng})=0,NA(),{form})"
    return form


def to_workbook(text, sheet_title="table", *, values=False):
    """Build a workbook from `.mdtbl` source.

    With `values=True`, computed cells and aggregates are `rowspec eval`'s
    values rather than formulas.

    Raises `Malformed` on a bad table, and `NotExpressible` on a valid one a
    spreadsheet cannot say.
    """
    cols, formulas, _rows, decls, _order, key = parse(text)
    rows, aggs = evaluate(text)
    folded = {}
    for c in cols:
        if c.casefold() in folded:
            raise NotExpressible(
                f"columns {folded[c.casefold()]!r} and {c!r} differ only by case, "
                f"and a spreadsheet table's header cannot hold both"
            )
        folded[c.casefold()] = c
    if not rows:
        raise NotExpressible("the table has no rows, and a spreadsheet table needs one")
    if values:
        col_formula, agg_formula = {}, {}
    else:
        col_formula = _column_formulas(cols, formulas, rows, key)
        agg_formula = {
            nm: _aggregate(fn, col, pred, cols, nm) for nm, (fn, col, pred) in decls.items()
        }

    wb = _Workbook()
    # A formula written here carries no cached value, so a spreadsheet must
    # compute every formula when the file opens. A values workbook holds none,
    # and openpyxl's default would still ask for a full recalculation.
    wb.calculation = CalcProperties(fullCalcOnLoad=not values)
    ws = wb.active
    _expressible(sheet_title)
    ws.title = sheet_title
    _append(ws, list(cols))
    for row in rows:
        # The KEY column is exported verbatim, never coerced. §9.16 makes a key
        # an identifier and §4.2 rule 6 compares identifiers as text, so `007`
        # and `7` are two rows -- §9.5 would refuse them if they were one.
        # Coercing them to a number merges two row identities into one, in the
        # export of a format whose whole unit of identity is the row id. §4.1.6
        # admits `007` at all precisely so zero-padded identifiers survive.
        _append(
            ws,
            [None if c in formulas else str(row[c]) if c == key else _value(row[c]) for c in cols],
        )
        for i, c in enumerate(cols, 1):
            if c in formulas:
                if values:
                    _static(ws.cell(ws.max_row, i), row[c])
                else:
                    _formula(ws.cell(ws.max_row, i), col_formula[c])
    table = Table(displayName=TABLE, ref=f"A1:{ws.cell(len(rows) + 1, len(cols)).coordinate}")
    table.tableColumns = [
        TableColumn(
            id=i,
            name=c,
            calculatedColumnFormula=TableFormula(attr_text=col_formula[c])
            if c in col_formula
            else None,
        )
        for i, c in enumerate(cols, 1)
    ]
    ws.add_table(table)
    for name in decls:
        _append(ws, [None])  # one blank row between the grid and the aggregates
        _append(ws, [name])
        if values:
            _static(ws.cell(ws.max_row, 2), aggs[name])
        else:
            _formula(ws.cell(ws.max_row, 2), agg_formula[name])
    return wb


def export_file(src, dst, sheet_title="table", *, values=False):
    with open(src, encoding="utf-8") as fh:
        wb = to_workbook(fh.read(), sheet_title, values=values)
    wb.save(dst)
    return dst
