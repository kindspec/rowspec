# SPDX-License-Identifier: Apache-2.0 OR MIT
"""The exported workbook, recalculated by LibreOffice, against `rowspec eval`.

kindspec/rowspec#36. A claim that a workbook recalculates, checked against our
own writer, proves nothing: `tests/test_xlsx_export.py` asserts what formulas
are written, and only a spreadsheet engine can say what they compute. So every
`.mdtbl` in the conformance tree that rowspec accepts and the exporter can
express is exported, recalculated in headless LibreOffice, and each computed
cell and aggregate is compared with `rowspec.table.evaluate` on the same file:

- a number must come back as the same number, compared at 15 significant
  digits, because LibreOffice saves no more than that -- a stored cell of
  `9007199254740992`, with no arithmetic involved, is saved as
  `9.00719925474099E+015` (#83);
- a `#REF!` must come back as a spreadsheet error value, never a number or a
  blank. Its name does not survive, and is not asserted.

The tree is used rather than tables written for this test because its cases
were written to break implementations, by authors who were not writing this
exporter.

LibreOffice's load-time recalculation is set to "always" in a throwaway
profile, so a cached value could never stand in for a computed one -- and the
exporter writes none. Every file goes through one `soffice` call.

`KNOWN` is the disagreements that are findings rather than bugs (#83). Each is
pinned to the value LibreOffice gives, so a new disagreement fails, and so does
one that goes away.

`just test` runs without the xlsx extra and skips this file. `just
test-libreoffice` sets `ROWSPEC_REQUIRE_LIBREOFFICE=1`, which turns a missing
openpyxl or `soffice` into a failure rather than a skip, and the `xlsx-extra`
CI job runs it.
"""

import glob
import os
import shutil
import subprocess

import pytest

REQUIRED = os.environ.get("ROWSPEC_REQUIRE_LIBREOFFICE") == "1"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
if REQUIRED:
    import openpyxl  # a skip here would be a check that cannot fail

    if SOFFICE is None:
        raise RuntimeError(
            "ROWSPEC_REQUIRE_LIBREOFFICE=1 says this test must RUN, and neither "
            "`soffice` nor `libreoffice` is on PATH, so it would have skipped."
        )
else:
    openpyxl = pytest.importorskip("openpyxl")
    if SOFFICE is None:
        pytest.skip("LibreOffice is not installed", allow_module_level=True)

from rowspec.table import Malformed, evaluate, parse  # noqa: E402
from rowspec_xlsx import NotExpressible, to_workbook  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# (case path, column or aggregate name): what LibreOffice gives. Every entry is
# §7's exact sum against the spreadsheet's left-to-right one (#83).
KNOWN = {
    ("merge/reordering-merge-keeps-the-sum/base.mdtbl", "s"): 0,
    ("merge/reordering-merge-keeps-the-sum/theirs.mdtbl", "s"): "#NUM!",
    ("eval/sum-huge-cancellation-positives-first/input.mdtbl", "s"): "#NUM!",
    ("eval/avg-finite-where-sum-overflows/input.mdtbl", "a"): "#NUM!",
}

# Properties the tree does not exercise, written for this test. A spreadsheet's
# `=` on text ignores case, and §4.2 rule 6 compares the characters.
EXTRA = {
    "extra/text-comparison-is-case-sensitive.mdtbl": (
        '| id | region | x = if(region = "EU", 1, 0) |\n| -- | -- | -: |\n'
        "| r_1 | EU |  |\n| r_2 | eu |  |\n\nkey := id\nn := sum(x)\n"
    ),
}

PROFILE = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load">
<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop>
</item>
</oor:items>
"""


def recalculate(paths, tmp):
    """Every file in one LibreOffice run; returns {path: recalculated path}."""
    os.makedirs(tmp / "profile" / "user")
    (tmp / "profile" / "user" / "registrymodifications.xcu").write_text(PROFILE)
    out = tmp / "recalculated"
    subprocess.run(
        [
            SOFFICE,
            f"-env:UserInstallation={(tmp / 'profile').as_uri()}",
            "--headless",
            "--norestore",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(out),
            *map(str, paths),
        ],
        check=True,
        capture_output=True,
        timeout=600,
    )
    got = {p: out / p.name for p in paths}
    missing = [str(p) for p, q in got.items() if not q.exists()]
    assert not missing, f"LibreOffice wrote no output for {missing}"
    return got


def agrees(want, cell):
    if isinstance(want, str) and want.startswith("#REF!"):
        return cell.data_type == "e"
    if not isinstance(cell.value, int | float) or isinstance(cell.value, bool):
        return False
    return float(f"{float(want):.15g}") == float(f"{float(cell.value):.15g}")


def test_libreoffice_recalculates_what_rowspec_eval_computes(tmp_path):
    sources = sorted(glob.glob(f"{ROOT}/conformance/cases/**/*.mdtbl", recursive=True))
    sources.append(f"{ROOT}/docs/example/parts.mdtbl")
    texts = {}
    for src in sources:
        with open(src, encoding="utf-8", newline="") as fh:
            texts[os.path.relpath(src, f"{ROOT}/conformance/cases")] = fh.read()
    texts.update(EXTRA)
    expected, refused = {}, 0
    (tmp_path / "exported").mkdir()
    for i, (case, text) in enumerate(texts.items()):
        try:
            rows, aggs = evaluate(text)
        except Malformed:
            continue
        try:
            wb = to_workbook(text)
        except NotExpressible:
            refused += 1
            continue
        dst = tmp_path / "exported" / f"{i:04d}.xlsx"
        wb.save(dst)
        expected[dst] = (case, parse(text)[1], rows, aggs)
    assert expected, "nothing was exported, so nothing would be compared"

    seen, wrong = {}, []
    for dst, back in recalculate(list(expected), tmp_path).items():
        case, formulas, rows, aggs = expected[dst]
        ws = openpyxl.load_workbook(back, data_only=True).active
        header = [c.value for c in ws[1]]
        cells = {}
        for j, col in enumerate(header, 1):
            if col in formulas:
                cells[col] = [(row[col], ws.cell(i, j)) for i, row in enumerate(rows, 2)]
        below = {ws.cell(r, 1).value: ws.cell(r, 2) for r in range(len(rows) + 2, ws.max_row + 1)}
        for name, want in aggs.items():
            cells.setdefault(name, []).append((want, below[name]))
        for name, pairs in cells.items():
            for want, cell in pairs:
                if (case, name) in KNOWN:
                    seen[(case, name)] = cell.value
                elif not agrees(want, cell):
                    wrong.append(f"{case} {name}: eval {want!r}, LibreOffice {cell.value!r}")
    assert not wrong, "\n".join(wrong)
    assert seen == KNOWN, f"a known disagreement changed: {seen} != {KNOWN}"
    print(f"{len(expected)} files compared, {refused} not expressible")
