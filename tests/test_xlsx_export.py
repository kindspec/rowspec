# SPDX-License-Identifier: Apache-2.0 OR MIT
"""The exporter, exercised only where its optional extra is installed.

`just test` runs with the extra ABSENT, because that is what `pip install
rowspec` produces and it is the configuration CI has to cover. These tests skip
there. A test that only ever skips is a check that cannot fail, so `just
test-xlsx` and the `xlsx-extra` CI job set `ROWSPEC_REQUIRE_XLSX=1` and the
import below stops being forgiving.

What is asserted here is the workbook's SHAPE: stored values reach the cells,
computed columns and aggregates arrive as formulas over structured references,
no formula holds a coordinate, and what a spreadsheet cannot say is refused.
Nothing here computes a formula, so nothing here claims what one evaluates to;
whether LibreOffice recalculates them to `rowspec eval`'s values is
kindspec/rowspec#36.

The `values=True` export (#89) holds no formula, so its check is here and not
there: every computed cell and aggregate it writes must equal `evaluate`'s,
exactly, over every conformance `.mdtbl` rowspec accepts.
"""

import glob
import os
import re
import subprocess
import sys
import threading
import zipfile

import pytest

if os.environ.get("ROWSPEC_REQUIRE_XLSX") == "1":
    try:
        import openpyxl
    except ModuleNotFoundError as exc:  # a skip here would be a check that cannot fail
        raise RuntimeError(
            "ROWSPEC_REQUIRE_XLSX=1 says these tests must RUN, and openpyxl is not "
            "installed, so they would have skipped instead — which is how a suite "
            "comes to report a pass over an exporter nothing executed. "
            "Run `just test-xlsx`, which installs the extra."
        ) from exc
else:
    openpyxl = pytest.importorskip("openpyxl")

from openpyxl.cell import _writer as cell_writer  # noqa: E402
from openpyxl.compat import safe_string as openpyxl_safe_string  # noqa: E402
from rowspec.table import Malformed, evaluate, parse  # noqa: E402
from rowspec_xlsx import NotExpressible, _static, export_file, to_workbook  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

TABLE = (
    "| id | item | qty | unit | total = qty * unit |\n"
    "| --- | --- | --: | --: | --: |\n"
    "| r_01 | widget | 10 | 12.00 |  |\n"
    "| r_02 | gadget | 3 | 4.50 |  |\n"
    "\n"
    "key := id\n"
    "grand := sum(total)\n"
)


def sheet(path):
    """Reload from disk. Cell TYPE is part of what is asserted here, and it is
    only real once the file has been written and read back."""
    wb = openpyxl.load_workbook(path)
    return wb[wb.sheetnames[0]]


def grid(path):
    return [list(r) for r in sheet(path).iter_rows(values_only=True)]


def test_the_header_and_the_stored_cells_arrive(tmp_path):
    src = tmp_path / "t.mdtbl"
    src.write_text(TABLE)
    rows = grid(export_file(str(src), str(tmp_path / "t.xlsx")))
    assert rows[0] == ["id", "item", "qty", "unit", "total"]
    assert rows[1][:4] == ["r_01", "widget", 10, 12.0]


GUARDED = "IF(ISNUMBER(rowspec[[#This Row],[{0}]]),rowspec[[#This Row],[{0}]],NA())"


def test_a_computed_column_arrives_as_a_structured_reference_formula(tmp_path):
    """Every operand is guarded, because a spreadsheet reads a blank as `0` and
    §8 reads it as `#REF!` -- the guard is what keeps the formula this format's."""
    src = tmp_path / "t.mdtbl"
    src.write_text(TABLE)
    path = export_file(str(src), str(tmp_path / "t.xlsx"))
    ws = sheet(path)
    want = "(" + GUARDED.format("qty") + "*" + GUARDED.format("unit") + ")"
    assert [ws.cell(r, 5).value for r in (2, 3)] == ["=" + want] * 2
    table = ws.tables["rowspec"]
    assert table.ref == "A1:E3"
    assert [c.name for c in table.tableColumns] == ["id", "item", "qty", "unit", "total"]
    assert table.tableColumns[4].calculatedColumnFormula.attr_text == want


def test_a_declared_aggregate_arrives_below_the_grid_as_subtotal(tmp_path):
    src = tmp_path / "t.mdtbl"
    src.write_text(TABLE)
    ws = sheet(export_file(str(src), str(tmp_path / "t.xlsx")))
    assert ws.cell(ws.max_row, 1).value == "grand"
    assert ws.cell(ws.max_row, 2).value == (
        "=IF(SUBTOTAL(103,rowspec[total])<>SUBTOTAL(102,rowspec[total]),"
        "NA(),SUBTOTAL(109,rowspec[total]))"
    )


def test_a_broken_total_arrives_as_a_formula_not_a_value(tmp_path):
    """§8's `#REF!` must not land in a spreadsheet as a number or a blank.

    A division by zero is decided by the data, so it arrives as the formula
    that will divide by zero. A name the header does not have is decided by the
    header, so that column arrives as an error outright -- a structured
    reference to an absent column is not a formula a spreadsheet will open.
    What either one evaluates to is asserted against LibreOffice (#36).
    """
    src = tmp_path / "t.mdtbl"
    src.write_text(
        "| id | a | z | t = a / z | u = nope + a |\n| -- | -: | -: | -: | -: |\n"
        "| r_1 | 5 | 0 |  |  |\n\nkey := id\n"
    )
    ws = sheet(export_file(str(src), str(tmp_path / "t.xlsx")))
    assert ws.cell(2, 4).value == "=(" + GUARDED.format("a") + "/" + GUARDED.format("z") + ")"
    assert ws.cell(2, 5).value == "=NA()"


# An A1 or R1C1 reference: column letters then a row number, or R<n>C<n>.
COORDINATE = re.compile(
    r"(?<![A-Za-z0-9_.])\$?[A-Za-z]{1,3}\$?[0-9]+(?![A-Za-z0-9_(])|\bR[0-9]*C[0-9]*\b"
)


def formulas_in(path):
    """Every formula in the saved file, read from the XML rather than through
    openpyxl, so a formula openpyxl would not surface is still checked."""
    out = []
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith(("xl/worksheets/", "xl/tables/")):
                xml = z.read(name).decode()
                out += re.findall(r"<f>(.*?)</f>", xml)
                out += re.findall(r"<calculatedColumnFormula>(.*?)</calculatedColumnFormula>", xml)
    return out


def test_no_formula_holds_a_coordinate(tmp_path):
    """#35: a coordinate anywhere in a formula is a bug.

    String literals and structured-reference brackets are removed first, so a
    column named `a1` or a literal `"B2"` is not mistaken for a coordinate.
    """
    src = tmp_path / "t.mdtbl"
    src.write_text(
        '| id | a1 | b | c = if(b > 1, a1 * 2, -b) | d = if(id <> "B2", c / 3, 0) |\n'
        "| -- | -: | -: | -: | -: |\n| r_1 | 5 | 2 |  |  |\n\nkey := id\n"
        "s := sum(c)\nn := count(d)\nlo := min(a1)\nhi := max(b)\nm := avg(c)\n"
    )
    found = formulas_in(export_file(str(src), str(tmp_path / "t.xlsx")))
    assert len(found) == 2 * 2 + 5, found  # two columns, cell and table, and five aggregates
    for f in found:
        bare = re.sub(r"\[[^\[\]]*\]", "", re.sub(r"\[[^\[\]]*\]", "", re.sub(r'"[^"]*"', "", f)))
        assert not COORDINATE.search(bare), f"coordinate in {f!r}"


def test_the_coordinate_check_can_fail():
    assert COORDINATE.search("SUM(A1:B2)")
    assert COORDINATE.search("$C$3*2")
    assert COORDINATE.search("R1C1")
    assert not COORDINATE.search("SUBTOTAL(109,rowspec)")


REFUSED_TABLES = [
    (
        "| id | d | a | c = cumulative(a) |\n| - | - | - | - |\n| r_1 | 1 | 2 |  |\n\n"
        "key := id\norder := by(d)\n",
        "preceding row",
    ),
    (
        "| id | g | a | t = sum(a where g = @g) |\n| - | - | - | - |\n| r_1 | x | 2 |  |\n\n"
        "key := id\n",
        "where",
    ),
    (
        '| id | g | a |\n| - | - | - |\n| r_1 | x | 2 |\n\nkey := id\ns := sum(a where g = "x")\n',
        "where",
    ),
    (
        '| id | q | t = if(q = "3", 1, 0) |\n| - | - | - |\n| r_1 | 3.0 |  |\n\nkey := id\n',
        "spelling",
    ),
    ("| id | ab | AB |\n| - | - | - |\n| r_1 | 1 | 2 |\n\nkey := id\n", "case"),
    ("| id | a |\n| - | - |\n\nkey := id\n", "no rows"),
    (
        "| id | a | t = a * 1.23456789012345678 |\n| - | - | - |\n| r_1 | 1 |  |\n\nkey := id\n",
        "significant digits",
    ),
    ("| id | a |\n| - | - |\n| r_1 | a\x0cb |\n\nkey := id\n", "no spreadsheet file"),
    ("| id | a |\n| - | - |\n| r_1 | a\x0bb |\n\nkey := id\n", "no spreadsheet file"),
    ("| id | a |\n| - | - |\n| r_1 | a\x00b |\n\nkey := id\n", "no spreadsheet file"),
    (
        '| id | g | t = if(g = "x\x0c", 1, 0) |\n| - | - | - |\n| r_1 | y |  |\n\nkey := id\n',
        "no spreadsheet file",
    ),
    ("| id | a |\n| - | - |\n| r_1 | a\ufffeb |\n\nkey := id\n", "no spreadsheet file"),
    (
        '| id | g | t = if(g = "x\uffff", 1, 0) |\n| - | - | - |\n| r_1 | y |  |\n\nkey := id\n',
        "no spreadsheet file",
    ),
]
REFUSED_IDS = [
    "row-relative",
    "group-where",
    "declaration-where",
    "text-vs-number",
    "case",
    "empty",
    "literal",
    "form-feed-cell",
    "vertical-tab-cell",
    "nul-cell",
    "form-feed-string-literal",
    "fffe-cell",
    "ffff-string-literal",
]
REFUSED = dict(zip(REFUSED_IDS, REFUSED_TABLES, strict=True))


@pytest.mark.parametrize(("table", "why"), REFUSED_TABLES, ids=REFUSED_IDS)
def test_what_a_spreadsheet_cannot_say_is_refused(table, why):
    with pytest.raises(NotExpressible, match=why):
        to_workbook(table)


def test_a_blank_test_on_a_number_column_is_exported(tmp_path):
    """`if(x = "", ...)` reads no spelling, so a number column is fine there:
    a number is never the empty string, to this format or to a spreadsheet."""
    wb = to_workbook(
        '| id | q | t = if(q = "", 0, q) |\n| - | - | - |\n'
        "| r_1 | 3.0 |  |\n| r_2 |  |  |\n\nkey := id\n"
    )
    assert wb.active.cell(2, 3).value.startswith('=IF(EXACT(rowspec[[#This Row],[q]],""),0,')


def test_a_stored_cell_beginning_with_equals_is_not_a_live_formula(tmp_path):
    """§5 makes a stored cell text. openpyxl reads a leading `=` as a formula.

    Exported without forcing the type, this cell arrives as `data_type == "f"`
    and survives save and reload as one — a computed cell the file never had,
    a `=A1` would smuggle a COORDINATE into the output, and a `.mdtbl` is a
    file anyone can open a pull request against.

    Asserted after a save and reload, not on the in-memory workbook, because
    the round trip is where the claim has to hold.
    """
    src = tmp_path / "t.mdtbl"
    src.write_text("| id | note |\n| --- | --- |\n| r_1 | =1+1 |\n\nkey := id\n")
    ws = sheet(export_file(str(src), str(tmp_path / "t.xlsx")))
    assert ws.cell(2, 2).value == "=1+1"
    assert ws.cell(2, 2).data_type == "s", "a stored cell was exported as a live formula"


def test_two_keys_that_differ_only_as_text_stay_two_rows(tmp_path):
    """§9.16 makes a key an identifier, not a number, and §4.2 rule 6 compares
    identifiers as text. `007` and `7` are two rows — §9.5 would refuse them if
    they were one — so an export that coerces the key column merges two row
    identities into one, in the export of a format named after the row id.

    §4.1.6 admits `007` at all precisely so zero-padded identifiers survive.
    """
    src = tmp_path / "t.mdtbl"
    src.write_text("| id | n |\n| --- | --: |\n| 007 | 1 |\n| 7 | 2 |\n\nkey := id\n")
    ws = sheet(export_file(str(src), str(tmp_path / "t.xlsx")))
    keys = [ws.cell(r, 1).value for r in (2, 3)]
    assert keys == ["007", "7"], keys
    assert [ws.cell(r, 1).data_type for r in (2, 3)] == ["s", "s"]
    # The VALUE column keeps the number policy: this is about the key alone.
    assert [ws.cell(r, 2).value for r in (2, 3)] == [1, 2]


def test_a_blank_stored_cell_arrives_as_a_cell_that_is_not_there(tmp_path):
    """`evaluate` returns a blank cell as `""`, which must not become `"None"`,
    `0`, or an EMPTY STRING cell.

    `data_type` is the assertion, not `value`, and that is the whole point of
    this test. Writing `""` emits `<c r="B2" t="inlineStr"/>` — a cell that
    exists and holds empty text, for which `ISBLANK()` is FALSE — while
    writing `None` emits no cell element at all. Measured on the saved
    sheet XML. openpyxl's reader returns `value is None` for BOTH, so a test
    asserting on the value passes over either and this one would be measuring
    nothing: `data_type` is `"n"` for the absent cell and `"inlineStr"` for the
    empty-string one.
    """
    src = tmp_path / "t.mdtbl"
    src.write_text("| id | note |\n| --- | --- |\n| r_1 |  |\n\nkey := id\n")
    ws = sheet(export_file(str(src), str(tmp_path / "t.xlsx")))
    assert ws.cell(2, 2).value is None
    assert ws.cell(2, 2).data_type == "n", "a blank cell was exported as empty TEXT"


def test_a_refused_table_is_refused_here_too(tmp_path):
    """Export defers to the parser rather than writing whatever it can."""
    with pytest.raises(Malformed):
        to_workbook("| id | n |\n| --- | --- |\n| r_01 | 1 |\n| r_01 | 2 |\n\nkey := id\n")


def test_the_module_entry_point_writes_a_file(tmp_path):
    dst = tmp_path / "parts.xlsx"
    env = dict(os.environ, PYTHONPATH=os.pathsep.join([f"{ROOT}/export", f"{ROOT}/reference"]))
    r = subprocess.run(
        [sys.executable, "-m", "rowspec_xlsx", f"{ROOT}/docs/example/parts.mdtbl", str(dst)],
        capture_output=True,
        text=True,
        env=env,
    )
    assert r.returncode == 0, r.stdout + r.stderr
    assert grid(str(dst))[0][0] == "id"


def test_the_module_entry_point_refuses_what_it_cannot_say(tmp_path):
    src = tmp_path / "t.mdtbl"
    src.write_text("| id | ab | AB |\n| - | - | - |\n| r_1 | 1 | 2 |\n\nkey := id\n")
    dst = tmp_path / "t.xlsx"
    env = dict(os.environ, PYTHONPATH=os.pathsep.join([f"{ROOT}/export", f"{ROOT}/reference"]))
    r = subprocess.run(
        [sys.executable, "-m", "rowspec_xlsx", str(src), str(dst)],
        capture_output=True,
        text=True,
        env=env,
    )
    assert r.returncode == 1, r.stdout + r.stderr
    assert "not exported" in r.stderr and "case" in r.stderr
    assert not dst.exists(), "a refused export wrote a file"


def test_tab_in_a_cell_is_exported():
    """Tab is the one C0 control a stored cell can hold that XML keeps, so the
    refusal of the others must not reach it."""
    wb = to_workbook("| id | a |\n| - | - |\n| r_1 | a\tb |\n\nkey := id\n")
    assert wb.active["B2"].value == "a\tb"


def test_a_sheet_title_xml_cannot_hold_is_refused():
    with pytest.raises(NotExpressible, match="no spreadsheet file"):
        to_workbook("| id | a |\n| - | - |\n| r_1 | 1 |\n\nkey := id\n", "a\x0cb")


# --- values=True (#89) -------------------------------------------------------

CASES = os.path.join(ROOT, "conformance", "cases")

# What the formula export refuses only because no formula says it. Values mode
# has `evaluate`'s value for each, so it exports them.
FOR_WANT_OF_A_FORMULA = [
    "row-relative",
    "group-where",
    "declaration-where",
    "text-vs-number",
    "literal",
    "form-feed-string-literal",
    "ffff-string-literal",
]


def test_values_mode_matches_rowspec_eval_on_every_accepted_case(tmp_path):
    """Every computed cell and every aggregate, read back from the saved file,
    is `evaluate`'s value: the same number exactly, an error cell for each
    `#REF!`, no cell for a blank. No cell anywhere is a formula.

    A refusal is allowed only for a reason values mode keeps (`case`, `no
    rows`, a character XML cannot hold), and every case the formula export
    refuses for want of a formula must be exported."""
    numbers = errors = blanks = 0
    refused, gained, wrong = [], 0, []
    for src in sorted(glob.glob(f"{CASES}/**/*.mdtbl", recursive=True)):
        case = os.path.relpath(src, CASES)
        with open(src, encoding="utf-8", newline="") as fh:
            text = fh.read()
        try:
            rows, aggs = evaluate(text)
        except Malformed:
            continue
        try:
            wb = to_workbook(text, values=True)
        except NotExpressible as e:
            refused.append(f"{case}: {e}")
            continue
        try:
            to_workbook(text)
        except NotExpressible:
            gained += 1
        dst = tmp_path / "out.xlsx"
        wb.save(dst)
        assert not formulas_in(dst), f"{case}: a formula in a values export"
        ws = sheet(dst)
        cols, formulas = parse(text)[:2]
        got = [
            (want, ws.cell(i, j), f"{c} row {i}")
            for i, row in enumerate(rows, 2)
            for j, c in enumerate(cols, 1)
            if c in formulas
            for want in [row[c]]
        ]
        below = {ws.cell(r, 1).value: ws.cell(r, 2) for r in range(len(rows) + 2, ws.max_row + 1)}
        got += [(want, below[name], name) for name, want in aggs.items()]
        for want, cell, where in got:
            if isinstance(want, str) and want.startswith("#REF!"):
                ok = cell.data_type == "e" and cell.value == "#REF!"
                errors += 1
            elif want == "":
                ok = cell.value is None and cell.data_type == "n"
                blanks += 1
            else:
                ok = type(cell.value) in (int, float) and cell.value == want
                ok = ok and repr(float(cell.value)) == repr(float(want))
                numbers += 1
            if not ok:
                wrong.append(f"{case} {where}: eval {want!r}, workbook {cell.value!r}")
    assert not wrong, "\n".join(wrong[:20])
    for r in refused:
        assert re.search("differ only by case|no rows|no spreadsheet file", r), r
    # A comparison over nothing is a check that cannot fail.
    assert numbers > 1000 and errors > 500 and blanks > 0 and gained > 100, (
        numbers,
        errors,
        blanks,
        gained,
    )


@pytest.mark.parametrize("name", FOR_WANT_OF_A_FORMULA)
def test_values_mode_exports_what_only_a_formula_could_not_say(name):
    table, _why = REFUSED[name]
    rows, aggs = evaluate(table)
    ws = to_workbook(table, values=True).active
    assert ws.max_row >= len(rows) + 1


@pytest.mark.parametrize("name", sorted(set(REFUSED) - set(FOR_WANT_OF_A_FORMULA)))
def test_values_mode_still_refuses_what_a_sheet_cannot_hold(name):
    table, why = REFUSED[name]
    with pytest.raises(NotExpressible, match=why):
        to_workbook(table, values=True)


def test_values_mode_writes_ref_as_an_error_cell_not_text(tmp_path):
    """A spreadsheet `SUM` skips text, so a `#REF!` written as text would let a
    total over it come out as a number. As an error cell it propagates."""
    src = tmp_path / "t.mdtbl"
    src.write_text(
        "| id | a | z | t = a / z |\n| -- | -: | -: | -: |\n"
        "| r_1 | 5 | 0 |  |\n| r_2 | 6 | 2 |  |\n\nkey := id\ns := sum(t)\n"
    )
    ws = sheet(export_file(str(src), str(tmp_path / "t.xlsx"), values=True))
    assert (ws.cell(2, 4).value, ws.cell(2, 4).data_type) == ("#REF!", "e")
    assert (ws.cell(3, 4).value, ws.cell(3, 4).data_type) == (3, "n")
    assert ws.cell(ws.max_row, 1).value == "s"
    assert (ws.cell(ws.max_row, 2).value, ws.cell(ws.max_row, 2).data_type) == ("#REF!", "e")


def test_values_mode_keeps_what_spreadsheet_arithmetic_loses(tmp_path):
    """The two disagreements the formula export has with `rowspec eval`, #83's
    and #88's: values mode writes eval's answer for both, to the last digit."""
    src = tmp_path / "t.mdtbl"
    src.write_text(
        "| id | a | b | k | c = a + b | d = if(c = 0.3, 1, 0) | z = a + b - k |\n"
        "| -- | -: | -: | -: | -: | -: | -: |\n"
        "| r_1 | 0.1 | 0.2 | 0.3 |  |  |  |\n"
        "| r_2 | 10000000000000000 | 0 | 0 |  |  |  |\n"
        "| r_3 | 1 | 0 | 0 |  |  |  |\n"
        "| r_4 | -10000000000000000 | 0 | 0 |  |  |  |\n\nkey := id\ns := sum(a)\n",
    )
    ws = sheet(export_file(str(src), str(tmp_path / "t.xlsx"), values=True))
    assert [ws.cell(2, c).value for c in (5, 6, 7)] == [
        0.30000000000000004,
        0,
        5.551115123125783e-17,
    ]
    assert ws.cell(ws.max_row, 2).value == 1.1


def test_the_default_export_is_still_formulas(tmp_path):
    src = tmp_path / "t.mdtbl"
    src.write_text(TABLE)
    ws = sheet(export_file(str(src), str(tmp_path / "t.xlsx")))
    assert ws.cell(2, 5).data_type == "f"


def test_the_module_entry_point_takes_values(tmp_path):
    src = tmp_path / "t.mdtbl"
    src.write_text(TABLE)
    dst = tmp_path / "t.xlsx"
    env = dict(os.environ, PYTHONPATH=os.pathsep.join([f"{ROOT}/export", f"{ROOT}/reference"]))
    r = subprocess.run(
        [sys.executable, "-m", "rowspec_xlsx", "--values", str(src), str(dst)],
        capture_output=True,
        text=True,
        env=env,
    )
    assert r.returncode == 0, r.stdout + r.stderr
    assert [c[4] for c in grid(str(dst))[1:3]] == [120, 13.5]
    assert not formulas_in(str(dst))


def accepted_cases():
    for src in sorted(glob.glob(f"{CASES}/**/*.mdtbl", recursive=True)):
        with open(src, encoding="utf-8", newline="") as fh:
            text = fh.read()
        try:
            yield os.path.relpath(src, CASES), text, evaluate(text)
        except Malformed:
            continue


def test_values_mode_cells_in_memory_are_eval_numbers():
    """The workbook `to_workbook` returns, before any save: a number cell holds
    a number, the one `evaluate` computed, never its spelling as text. The
    exact digits are a matter for the save, not for the cell."""
    seen, wrong = 0, []
    for case, text, (rows, aggs) in accepted_cases():
        try:
            ws = to_workbook(text, values=True).active
        except NotExpressible:
            continue
        cols, formulas = parse(text)[:2]
        pairs = [
            (f"{c} row {i}", row[c], ws.cell(i, j).value)
            for i, row in enumerate(rows, 2)
            for j, c in enumerate(cols, 1)
            if c in formulas
        ]
        below = {
            ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(len(rows) + 2, ws.max_row + 1)
        }
        pairs += [(name, want, below[name]) for name, want in aggs.items()]
        for where, want, got in pairs:
            if isinstance(want, int | float):
                seen += 1
                if not (isinstance(got, int | float) and repr(got) == repr(want)):
                    wrong.append(f"{case} {where}: eval {want!r}, cell {got!r}")
    assert not wrong, "\n".join(wrong[:20])
    assert seen > 500, seen


STORED_17 = (
    "| id | a | t = a * 1 |\n| -- | -: | -: |\n| r_1 | 0.30000000000000004 |  |\n\nkey := id\n"
)


@pytest.mark.parametrize("values", [False, True], ids=["formulas", "values"])
def test_a_stored_number_keeps_its_17th_digit(tmp_path, values):
    """#269: openpyxl alone saves a number with 16 significant digits, so a
    stored `0.30000000000000004` came back as `0.3`, in either mode."""
    src = tmp_path / "t.mdtbl"
    src.write_text(STORED_17)
    ws = sheet(export_file(str(src), str(tmp_path / "t.xlsx"), values=values))
    assert repr(ws.cell(2, 2).value) == "0.30000000000000004"


def test_the_exact_digits_hook_still_applies(tmp_path):
    """The exact spelling depends on replacing `openpyxl.cell._writer.safe_string`
    during a save, a private name. If openpyxl stops writing cells through it,
    the replacement silently does nothing, and this fails on the saved XML.
    It also checks that the save puts openpyxl's own function back."""
    wb = to_workbook(
        "| id | a | u = 0 * -1 |\n| -- | -: | -: |\n"
        "| r_1 | 0.30000000000000004 |  |\n\nkey := id\n",
        values=True,
    )
    dst = tmp_path / "t.xlsx"
    wb.save(dst)
    with zipfile.ZipFile(dst) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode()
    assert "<v>0.30000000000000004</v>" in xml and "<v>-0.0</v>" in xml, xml
    # Against openpyxl's own function, not a copy taken here: an earlier save
    # that failed to restore it would have left the replacement in place.
    assert cell_writer.safe_string is openpyxl_safe_string


def test_minus_zero_is_written_as_minus_zero(tmp_path):
    """No conformance case computes -0.0, so this one does: `0 * -1` is -0.0
    under binary64, and `%.16g` spells it `-0`, which reads back as the
    integer 0."""
    src = tmp_path / "t.mdtbl"
    src.write_text("| id | a | u = a * -1 |\n| -- | -: | -: |\n| r_1 | 0 |  |\n\nkey := id\n")
    assert repr(evaluate(src.read_text())[0][0]["u"]) == "-0.0"
    ws = sheet(export_file(str(src), str(tmp_path / "t.xlsx"), values=True))
    assert repr(ws.cell(2, 3).value) == "-0.0"


def test_to_workbook_defaults_to_formulas():
    wb = to_workbook(TABLE)
    assert wb.active.cell(2, 5).data_type == "f"
    assert wb.calculation.fullCalcOnLoad is True


def test_values_mode_asks_for_no_recalculation(tmp_path):
    """A values workbook has no formula, and openpyxl's default `calcPr` would
    still ask the spreadsheet for a full recalculation on load."""
    src = tmp_path / "t.mdtbl"
    src.write_text(TABLE)
    path = export_file(str(src), str(tmp_path / "t.xlsx"), values=True)
    assert not openpyxl.load_workbook(path).calculation.fullCalcOnLoad
    path = export_file(str(src), str(tmp_path / "f.xlsx"))
    assert openpyxl.load_workbook(path).calculation.fullCalcOnLoad


@pytest.mark.parametrize("v", ["hello", True, None, [1]])
def test_a_computed_value_that_is_not_a_number_blank_or_ref_is_refused(v):
    cell = openpyxl.Workbook().active["A1"]
    with pytest.raises(TypeError, match="a number, a blank or a #REF!"):
        _static(cell, v)


@pytest.mark.parametrize(
    ("argv", "code"),
    [(["--val", "a.mdtbl", "b.xlsx"], 2), (["-h"], 0), (["only-one"], 2)],
    ids=["abbreviation", "help", "one-path"],
)
def test_the_module_entry_point_parses_its_arguments(argv, code, capsys):
    from rowspec_xlsx.__main__ import main

    assert main(argv) == code


def test_concurrent_saves_are_each_exact(tmp_path):
    """`_Workbook.save` swaps a module-level function. Unserialized, one
    thread's restore lands inside another thread's save, and the rest of that
    file is written with 16 digits. A tiny switch interval makes threads
    interleave inside a save, which is where that shows. The race is a matter
    of chance: without the lock this test failed 10 runs in 10, measured, and
    a single pass of it would not prove the lock is there."""
    old = sys.getswitchinterval()
    sys.setswitchinterval(1e-6)
    text = (
        "| id | a |\n| -- | -: |\n"
        + "".join(f"| r_{i} | 0.30000000000000004 |\n" for i in range(500))
        + "\nkey := id\n"
    )
    books = [to_workbook(text, values=True) for _ in range(8)]
    errors = []

    def save(wb, dst):
        try:
            wb.save(dst)
        except Exception as e:  # a thread's exception is otherwise lost
            errors.append(e)

    try:
        for trial in range(20):
            paths = [tmp_path / f"{trial}-{i}.xlsx" for i in range(len(books))]
            threads = [
                threading.Thread(target=save, args=a) for a in zip(books, paths, strict=True)
            ]
            for t in threads:
                t.start()
            for t in threads:
                t.join()
            for p in paths:
                with zipfile.ZipFile(p) as z:
                    xml = z.read("xl/worksheets/sheet1.xml").decode()
                assert xml.count("<v>0.30000000000000004</v>") == 500, p.name
    finally:
        sys.setswitchinterval(old)
    assert not errors, errors
    assert cell_writer.safe_string is openpyxl_safe_string
